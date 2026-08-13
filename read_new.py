import sys, io, os, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Copy files to /tmp with simple names to avoid encoding issues
src_dir = '/docs'
files = os.listdir(src_dir)
print("Files:", files)

mapping = {}
for i, f in enumerate(files):
    ext = os.path.splitext(f)[1]
    dst = f'/tmp/file{i}{ext}'
    shutil.copy(os.path.join(src_dir, f), dst)
    mapping[i] = (f, dst)

for i, (orig, path) in mapping.items():
    print(f"\n===== {orig} =====")
    if path.endswith('.docx'):
        try:
            from docx import Document
            d = Document(path)
            for p in d.paragraphs:
                if p.text.strip():
                    print(p.text)
            for t in d.tables:
                for row in t.rows:
                    print(" | ".join(c.text.strip() for c in row.cells))
        except Exception as e:
            print(f"ERROR: {e}")
    elif path.endswith('.pdf'):
        # Check header
        with open(path, 'rb') as f:
            header = f.read(8)
        print(f"Header bytes: {header}")
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            print(f"Pages: {len(reader.pages)}")
            for j, page in enumerate(reader.pages):
                print(f"--- Page {j+1} ---")
                print(page.extract_text())
        except Exception as e:
            print(f"pypdf ERROR: {e}")
        try:
            import subprocess
            subprocess.run(['pip', 'install', 'pdfplumber', '-q'], capture_output=True)
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                print(f"pdfplumber Pages: {len(pdf.pages)}")
                for j, page in enumerate(pdf.pages):
                    print(f"--- Page {j+1} ---")
                    print(page.extract_text())
        except Exception as e:
            print(f"pdfplumber ERROR: {e}")
