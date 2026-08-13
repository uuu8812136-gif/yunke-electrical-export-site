from docx import Document
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Read PPT docx
print("===== 云科简介(2).docx - 认证相关内容 =====")
d = Document('/work/云科简介(2).docx')
full_text = []
for p in d.paragraphs:
    if p.text.strip():
        full_text.append(p.text)
for t in d.tables:
    for row in t.rows:
        full_text.append(" | ".join(c.text.strip() for c in row.cells))

# Search for certification keywords
keywords = ['CE', 'CCC', 'ISO', 'IEC', '认证', '证书', 'certification', 'Certification', 'certif', 'standard', 'Standard', '标准', 'SASO', 'PLN', 'SNI', 'GB', 'IEC 61439', 'IEC 62271']
print("\n--- 包含认证/标准关键词的段落 ---")
seen = set()
for line in full_text:
    if any(kw in line for kw in keywords):
        if line not in seen:
            print(f">>> {line}")
            print()
            seen.add(line)

# Count occurrences
print("\n--- 关键词出现次数 ---")
text = "\n".join(full_text)
for kw in ['CE', 'CCC', 'ISO 9001', 'ISO 14001', 'ISO 45001', 'ISO', 'IEC 61439', 'IEC 62271', 'IEC', 'SASO', 'PLN', 'SNI', 'GB']:
    count = text.count(kw)
    print(f"{kw}: {count} 次")
