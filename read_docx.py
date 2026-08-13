from docx import Document
import sys

for path in ['云科简介(2).docx', '同行对标企业网站及优缺点.docx']:
    print(f"\n===== {path} =====")
    try:
        d = Document(path)
        for p in d.paragraphs:
            if p.text.strip():
                print(p.text)
        for t in d.tables:
            for row in t.rows:
                print(" | ".join(c.text.strip() for c in row.cells))
    except Exception as e:
        print(f"ERROR: {e}")
