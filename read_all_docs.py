from docx import Document
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

for label, path in [('同行对标', '/tmp/a.docx'), ('云科简介', '/tmp/b.docx')]:
    print(f"\n===== {label} =====")
    try:
        d = Document(path)
        for p in d.paragraphs:
            if p.text.strip():
                print(p.text)
        for t in d.tables:
            for row in t.rows:
                print(" | ".join(c.text.strip() for c in row.cells))
    except Exception as e:
        print(f"ERROR: {e}")

print("\n===== PDF =====")
try:
    from pypdf import PdfReader
    reader = PdfReader('/tmp/c.pdf')
    for i, page in enumerate(reader.pages):
        print(f"--- Page {i+1} ---")
        print(page.extract_text())
except Exception as e:
    print(f"ERROR: {e}")
