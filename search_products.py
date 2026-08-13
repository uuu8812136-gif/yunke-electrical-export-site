from docx import Document
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Read PPT docx
print("===== 云科简介(2).docx =====")
d = Document('/work/云科简介(2).docx')
full_text = []
for p in d.paragraphs:
    if p.text.strip():
        full_text.append(p.text)
for t in d.tables:
    for row in t.rows:
        full_text.append(" | ".join(c.text.strip() for c in row.cells))

text = "\n".join(full_text)

# Search for product-related keywords
keywords = ['产品', '热销', 'Product', 'product', 'Main product', '主要产品', '系列', 'Series', 'series', 'switchgear', 'transformer', 'substation', '开关柜', '变压器', '箱变', '新能源']
print("\n--- 包含产品相关关键词的段落 ---")
for line in full_text:
    if any(kw in line for kw in keywords):
        print(f">>> {line}")
        print()
